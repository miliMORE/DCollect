import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from django.utils import timezone
from django.utils.text import slugify

from catalog.models import ComparatorOrganization, SiteSettings
from collection.models import CountyCase, IndividualResponse
from collection.services import ensure_case_structure, iter_case_document_files

from .analysis import build_analysis_rows, orgs_with_pay, rollup_rows


NAVY = "1B2A4A"
GOLD = "C6A15B"
INK = "1F2933"
MUTED = "E8E1D4"
WHITE = "FFFFFF"
THIN = Border(
    left=Side(style="thin", color="C5BBA8"),
    right=Side(style="thin", color="C5BBA8"),
    top=Side(style="thin", color="C5BBA8"),
    bottom=Side(style="thin", color="C5BBA8"),
)


def _font(bold=False, color=INK, size=11):
    return Font(name="Calibri", bold=bold, color=color, size=size)


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def _header(ws, title, last_col="G"):
    ws.merge_cells(f"A1:{last_col}1")
    cell = ws["A1"]
    cell.value = title
    cell.font = _font(True, WHITE, 14)
    cell.fill = _fill(NAVY)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28


def _write_heads(ws, heads, row=4):
    for i, h in enumerate(heads, 1):
        cell = ws.cell(row, i, h)
        cell.font = _font(True, WHITE)
        cell.fill = _fill(GOLD)
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _strip_illegal_excel(text):
    try:
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

        return ILLEGAL_CHARACTERS_RE.sub("", text)
    except Exception:
        return "".join(ch for ch in text if ord(ch) >= 32 or ch in "\t\n")


def _cell_value(value):
    if value is None:
        return ""
    if hasattr(value, "quantize"):
        return float(value)
    if isinstance(value, str):
        return _strip_illegal_excel(value)
    return value


def _sheet_title(name, fallback="Sheet"):
    cleaned = "".join(ch for ch in (name or fallback) if ch not in r"[]:*?/\\")
    cleaned = cleaned.strip() or fallback
    return cleaned[:31]


def _set_widths(ws, widths):
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _label_value(ws, row, label, value):
    ws.cell(row, 1, label).font = _font(True)
    ws.cell(row, 2, _cell_value(value) if value not in (None, "") else "")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)


def _src_equivalent(row):
    code = getattr(row, "src_equivalent_code", None)
    if callable(code):
        value = code()
        if value:
            return value
    value = getattr(row, "src_equivalent", None)
    if isinstance(value, str):
        return value
    if getattr(row, "src_equivalent_id", None) and value:
        return value.code
    return ""


INDIVIDUAL_HEADS = [
    "County",
    "Submitted",
    "Job title",
    "Benchmark match",
    "Job group",
    "Department",
    "Terms",
    "Years in role",
    "Basic",
    "House",
    "Commuter",
    "Airtime",
    "Other",
    "Total cash",
    "Name",
    "Email",
]


def _individual_row_values(row):
    group = ""
    if row.job_id and getattr(row.job, "job_group_id", None):
        group = row.job.job_group.code
    return [
        row.county.name,
        row.submitted_at.strftime("%Y-%m-%d %H:%M") if row.submitted_at else "",
        row.job_title,
        row.job.title if row.job_id else "",
        group,
        row.department,
        row.get_employment_type_display(),
        float(row.years_in_role) if row.years_in_role is not None else "",
        float(row.current_basic) if row.current_basic is not None else "",
        float(row.house_allowance) if row.house_allowance is not None else "",
        float(row.commuter_allowance) if row.commuter_allowance is not None else "",
        float(row.airtime) if row.airtime is not None else "",
        float(row.other_allowances) if row.other_allowances is not None else "",
        float(row.total_cash()),
        row.respondent_name or "",
        row.respondent_email or "",
    ]


def _fill_individuals_sheet(ws, heading, responses):
    last = get_column_letter(len(INDIVIDUAL_HEADS))
    _header(ws, heading, last)
    _write_heads(ws, INDIVIDUAL_HEADS)
    r = 5
    for row in responses:
        for i, value in enumerate(_individual_row_values(row), 1):
            ws.cell(r, i, _cell_value(value) if not isinstance(value, float) else value)
        r += 1
    _set_widths(ws, [18, 18, 28, 28, 12, 22, 16, 12, 12, 12, 12, 12, 12, 14, 22, 28])
    return ws


def _write_job_analysis_sheet(wb, sheet_name, heading, analysis_rows, orgs, note=""):
    aheads = [
        "Job Group",
        "SRC Equivalent",
        "Official",
        "Staff n",
        "County P25",
        "County P50",
        "County P75",
    ] + [f"{org.name} median" for org in orgs] + [
        "Market n",
        "Market P25",
        "Market P50",
        "Market P75",
        "Ratio",
        "Position",
        "Scale",
    ]
    last_col = get_column_letter(max(len(aheads), 1))
    wa = wb.create_sheet(_sheet_title(sheet_name))
    _header(wa, heading, last_col)
    _write_heads(wa, aheads)
    r = 5
    for item in analysis_rows:
        if not item.get("has_county"):
            continue
        values = [
            item.get("job_group") or "",
            item.get("src_equivalent") or "",
            _cell_value(item.get("official")),
            item.get("staff_n") or 0,
            _cell_value(item.get("county_p25")),
            _cell_value(item.get("county")),
            _cell_value(item.get("county_p75")),
        ]
        comps = list(item.get("comparators") or [])
        for offset, _org in enumerate(orgs):
            value = comps[offset] if offset < len(comps) else None
            values.append(_cell_value(value))
        values.extend(
            [
                item.get("market_n") or 0,
                _cell_value(item.get("market_p25")),
                _cell_value(item.get("market")),
                _cell_value(item.get("market_p75")),
                _cell_value(item.get("position")),
                item.get("position_label") or "",
                item.get("scale_label") or "",
            ]
        )
        for i, value in enumerate(values, 1):
            wa.cell(r, i, value)
        r += 1
    if note:
        note_row = max(r + 1, 8)
        wa.cell(note_row, 1, "Note")
        note_end = max(len(aheads), 2)
        wa.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=note_end)
        wa.cell(note_row, 2, note)
    pack_widths = [14, 16, 14, 10, 14, 14, 14] + [24] * len(orgs) + [12, 14, 14, 14, 10, 16, 22]
    _set_widths(wa, pack_widths)
    return wa


def _write_comparator_columns_sheet(wb, sheet_name, heading, tagged_rows, orgs):
    heads = ["County", "Status", "Job Group", "SRC Equivalent"] + [org.name for org in orgs]
    last_col = get_column_letter(max(len(heads), 1))
    ws = wb.create_sheet(_sheet_title(sheet_name))
    _header(ws, heading, last_col)
    _write_heads(ws, heads)
    r = 5
    for item in tagged_rows:
        if not item.get("has_county"):
            continue
        values = [
            item.get("county_name") or "",
            item.get("county_status") or "",
            item.get("job_group") or "",
            item.get("src_equivalent") or "",
        ]
        comps = list(item.get("comparators") or [])
        for offset, _org in enumerate(orgs):
            value = comps[offset] if offset < len(comps) else None
            values.append(_cell_value(value))
        for i, value in enumerate(values, 1):
            ws.cell(r, i, value)
        r += 1
    _set_widths(ws, [20, 16, 12, 16] + [24] * len(orgs))
    return ws


def _write_rollup_sheet(wb, title, heading, rollups, last_col="M"):
    ws = wb.create_sheet(_sheet_title(title))
    _header(ws, heading, last_col)
    heads = [
        title,
        "Jobs with pay",
        "Average",
        "County P25",
        "County P50",
        "County P75",
        "Market n",
        "Market P25",
        "Market P50",
        "Market P75",
        "County P50 / market",
        "Position",
        "Market scale",
    ]
    _write_heads(ws, heads)
    r = 5
    for item in rollups:
        values = [
            item["label"],
            item["jobs_with_pay"],
            _cell_value(item.get("average")),
            _cell_value(item.get("county_p25")),
            _cell_value(item.get("county")),
            _cell_value(item.get("county_p75")),
            item.get("market_n") or 0,
            _cell_value(item.get("market_p25")),
            _cell_value(item.get("market")),
            _cell_value(item.get("market_p75")),
            _cell_value(item.get("position")),
            item.get("position_label") or "",
            item.get("scale_label") or "",
        ]
        for i, value in enumerate(values, 1):
            ws.cell(r, i, value)
        r += 1
    _set_widths(ws, [18, 14, 14, 14, 14, 14, 12, 14, 14, 14, 18, 16, 22])
    return ws


def build_county_pack(case: CountyCase, include_market=True) -> bytes:
    ensure_case_structure(case)
    settings_obj = SiteSettings.load()
    q = case.get_questionnaire()
    if q is None:
        from collection.models import Questionnaire

        q = Questionnaire(case=case)
    wb = Workbook()

    ws = wb.active
    ws.title = "Instructions"
    _header(ws, f"{settings_obj.survey_title}  |  FIELD-READY DATA COLLECTION TOOL", "B")
    ws["A2"] = "Purpose"
    ws["B2"] = _cell_value(settings_obj.purpose_text)
    ws["A3"] = "How to use"
    ws["B3"] = _cell_value(settings_obj.how_to_use_text)
    ws["A4"] = "Confidentiality"
    ws["B4"] = _cell_value(settings_obj.confidentiality_text)
    ws["A5"] = "County file"
    ws["B5"] = f"{case.county.name} — {case.survey_period.name} — {case.get_status_display()}"
    for r in range(2, 6):
        ws[f"A{r}"].font = _font(True)
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 48
    _set_widths(ws, [22, 110])

    wq = wb.create_sheet("County Questionnaire")
    _header(wq, "COUNTY GOVERNMENT SALARY & REMUNERATION QUESTIONNAIRE", "F")
    wq["B2"] = "Complete one questionnaire per County Government."
    rows = [
        ("SECTION A: INSTITUTION & RESPONDENT INFORMATION", ""),
        ("County :", case.county.name),
        ("Department / Directorate / Unit", q.department),
        ("Respondent Name (Confidential For Consultant's Use only)", q.respondent_name),
        ("Designation (Confidential For Consultant's Use only)", q.respondent_designation),
        ("Telephone (For Consultant's use only)", q.respondent_telephone),
        ("Email (For Consultant's use only)", q.respondent_email),
        ("Survey Period:", case.survey_period.name),
        ("SECTION B: WORKFORCE PROFILE", ""),
        ("No Permanent & Pensionable Employees", q.permanent_pensionable),
        ("No of Contract Employees", q.contract_employees),
        ("No of Temporary/Casual Employees", q.temporary_casual),
        ("Other Employees", q.other_employees),
        ("Total Employees", q.total_employees()),
        ("SECTION C: REMUNERATION POLICY & PHILOSOPHY", ""),
        ("Do you have an Approved Remuneration/Compensation Policy?", q.approved_policy),
        ("Is your Compensation philosophy formally defined?", q.philosophy_defined),
        ("What is your market position?", q.market_position),
        ("Which Year was remuneration structure last reviewed", q.last_review_year),
        ("What were the main triggers for salary review eg Inflation", q.review_triggers),
        ("SECTION D: SALARY STRUCTURE", ""),
        ("Number of salary grades in your organization", q.number_of_grades),
        ("Factors contributing to Salary progression eg Annual Increment", q.progression_factors),
        ("If Other, specify the factors contributing to salary progression", q.progression_factors_notes),
        ("Are there situations when Employees are paid above maximum?", q.paid_above_maximum),
        ("If Yes, why?", q.paid_above_maximum_why),
        ("SECTION E: ALLOWANCES & BENEFITS", ""),
        ("Main allowances — Remunerative Allowances (Paid monthly)", q.remunerative_allowances),
        ("Main allowances — Non-Remunerative allowances", q.non_remunerative_allowances),
        ("Approx. annual allowance expenditure (KES)", q.annual_allowance_expenditure),
        ("Medical insurance provided?", q.medical_insurance),
        ("How is medical cover arranged for employees?", q.medical_cover_how),
        ("Explain how medical insurance is handled", q.medical_cover_notes),
        ("Annual employer medical cost per employee (KES)", q.medical_cost_per_employee),
        ("Pension employer contribution (%)", q.pension_employer_pct),
        ("Pension employee contribution (%)", q.pension_employee_pct),
        ("Gratuity (%):", q.gratuity_pct),
        ("Other significant benefits", q.other_significant_benefits),
        ("SECTION F: RECRUITMENT, RETENTION & MARKET COMPETITIVENESS", ""),
        ("Challenges affecting Recruitment", q.recruitment_challenges),
        ("Job families with greatest recruitment challenges (Difficult to get the right candidates)", q.job_families_recruitment),
        ("Job families with greatest retention challenges", q.job_families_retention),
        ("Main reasons for recruitment challenges", q.challenge_reasons),
        ("Main reasons for retention challenges", q.retention_challenge_reasons),
        ("External salary benchmarking frequency ie After how long do you do external salary benchmarking?", q.benchmarking_frequency),
        ("Usual comparator organizations", q.usual_comparators),
        ("SECTION G: INTERNAL & EXTERNAL EQUITY", ""),
        ("Overall internal equity assessment (Equal pay for work of equal value)", q.internal_equity),
        ("Are there Significant pay differences for similar work?", q.pay_differences),
        ("Main causes of pay differences (If Yes in the above, what are the causes of pay differentials?)", q.pay_difference_causes),
        ("SECTION H: PAY MIX & PROGRESSION", ""),
        ("Basic salary as % of cash remuneration ((Basic Salary ÷ Total Cash Remuneration) × 100)", q.basic_pct),
        ("Allowances as % of cash remuneration", q.allowances_pct),
        ("Other cash payments as % ((Other Cash Payments ÷ Total Cash Payments) x100)", q.other_cash_pct),
        ("Annual salary progression method", q.progression_method),
        ("Average years in a salary grade", q.avg_years_in_grade),
        ("What Factors do determining salary placement?", q.placement_factors),
        ("SECTION I: AFFORDABILITY & IMPLEMENTATION", ""),
        ("Personnel/remuneration expenditure as % of budget", q.personnel_pct_budget),
        ("Can County accommodate general remuneration increase of about 15%?", q.can_accommodate_increase),
        ("Most sustainable implementation approach", q.implementation_approach),
        ("Preferred implementation period", q.implementation_period),
        ("Given budget constraints, which category of employees would you prioritize for salary increase? and why?", q.priority_categories),
        ("SECTION J: DOCUMENTS & KEY OBSERVATIONS", ""),
        ("Current salary structure attached?", q.salary_structure_attached),
        ("Allowance schedule attached?", q.allowance_schedule_attached),
        ("Benefits schedule attached?", q.benefits_schedule_attached),
        ("Previous salary survey attached?", q.previous_survey_attached),
        ("Three key remuneration challenges", q.three_challenges),
        ("Three recommended changes", q.three_recommendations),
        ("Other issues/comments", q.other_comments),
        ("SECTION K: RESPONDENT CERTIFICATION", ""),
        ("Name", case.certified_name),
        ("Designation", case.certified_designation),
        ("Signature/Initials", case.certified_initials),
        ("Date", case.certified_at.strftime("%Y-%m-%d") if case.certified_at else ""),
    ]
    if q and q.pk:
        from catalog.models import QuestionItem

        extras = QuestionItem.objects.filter(is_builtin=False, is_active=True).order_by("section", "sort_order")
        extra_rows = []
        for item in extras:
            answer = q.extra_answers.filter(question=item).first()
            extra_rows.append((item.label, answer.value if answer else ""))
        if extra_rows:
            rows.append(("ADDITIONAL QUESTIONS", ""))
            rows.extend(extra_rows)
    r = 4
    for label, value in rows:
        if label.startswith("SECTION"):
            wq.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            cell = wq.cell(r, 1, label)
            cell.font = _font(True, WHITE)
            cell.fill = _fill(NAVY)
        else:
            _label_value(wq, r, label, value)
        r += 1
    _set_widths(wq, [70, 28, 16, 16, 16, 16])

    wb_jobs = wb.create_sheet("Benchmark Positions")
    _header(wb_jobs, "JOB GROUP PAY DATA IN COUNTY GOVERNMENT", "I")
    headers = [
        "Job Group",
        "SRC Equivalent",
        "Minimum Basic Salary",
        "Maximum Basic Salary",
        "Midpoint (analysis)",
        "House Allowance",
        "Commuter Allowance",
        "Air Time",
        "Total of Other Allowances",
    ]
    for i, h in enumerate(headers, 1):
        cell = wb_jobs.cell(4, i, h)
        cell.font = _font(True, WHITE)
        cell.fill = _fill(GOLD)
    r = 5
    from catalog.models import JobGroup as _JobGroup

    for row in case.job_rows.select_related("job_group").filter(
        job_group__code__in=_JobGroup.PAY_CODES
    ):
        if row.not_applicable:
            values = [row.job_group.title, _src_equivalent(row), "N/A", "", "", "", "", "", ""]
        else:
            values = [
                row.job_group.title,
                _src_equivalent(row),
                row.min_basic,
                row.max_basic,
                row.midpoint_basic(),
                row.house_allowance,
                row.commuter_allowance,
                row.airtime,
                row.other_allowances,
            ]
        for i, value in enumerate(values, 1):
            wb_jobs.cell(r, i, _cell_value(value))
        r += 1
    _set_widths(wb_jobs, [22, 16, 18, 18, 16, 16, 16, 12, 24])

    staff_rows = IndividualResponse.objects.filter(
        county=case.county, survey_period=case.survey_period
    ).select_related("county", "job", "job__job_group").order_by("-submitted_at")
    wi = wb.create_sheet("Individual responses")
    _fill_individuals_sheet(
        wi,
        f"STAFF SURVEY RESPONSES — {case.county.name}",
        staff_rows,
    )

    orgs = orgs_with_pay(
        ComparatorOrganization.objects.filter(is_active=True, include_in_market_median=True).order_by(
            "sort_order", "name"
        ),
        case.survey_period,
    )
    if include_market:
        if not orgs:
            wc = wb.create_sheet("Comparator")
            _header(wc, "COMPARATOR 1 – SALARY & BENEFITS DATA", "K")
        for index, org in enumerate(orgs, 1):
            title = "Comparator" if index == 1 else f"Comparator {index}"
            wc = wb.create_sheet(_sheet_title(title))
            _header(wc, f"COMPARATOR {index} – {org.name.upper()} SALARY & BENEFITS DATA", "K")
            comparator_heads = [
                "Job Title",
                "Job Group",
                "Minimum Basic Salary",
                "Maximum Basic Salary",
                "Median Basic Salary (analysis)",
                "House Allowance",
                "Commuter Allowance",
                "Air Time",
                "Total of Other Allowances",
            ]
            for i, h in enumerate(comparator_heads, 1):
                cell = wc.cell(4, i, h)
                cell.font = _font(True, WHITE)
                cell.fill = _fill(GOLD)
            r = 5
            for row in org.pay_rows.filter(survey_period=case.survey_period).select_related(
                "job", "job__job_group"
            ):
                values = [row.job.title, getattr(row.job.job_group, "code", "")]
                if row and not row.not_applicable:
                    values += [
                        row.min_basic,
                        row.max_basic,
                        row.midpoint_basic(),
                        row.house_allowance,
                        row.commuter_allowance,
                        row.airtime,
                        row.other_allowances,
                    ]
                else:
                    values += ["", "", "", "", "", "", ""]
                for i, value in enumerate(values, 1):
                    wc.cell(r, i, _cell_value(value))
                r += 1
            _set_widths(wc, [36, 12, 18, 18, 16, 16, 16, 12, 24])

    wm = wb.create_sheet("Benefits Matrix")
    _header(wm, "EMPLOYEE BENEFITS COMPARISON MATRIX", "G")
    bheads = ["Benefit", "Provided?", "Interest charged?", "Employer Contribution", "Employee Contribution", "Eligibility criteria", "Explanatory comments"]
    for i, h in enumerate(bheads, 1):
        cell = wm.cell(4, i, h)
        cell.font = _font(True, WHITE)
        cell.fill = _fill(GOLD)
    r = 5
    for row in case.benefit_rows.select_related("benefit"):
        values = [row.benefit.name, row.provided, row.interest_charged, row.employer_contribution, row.employee_contribution, row.eligibility, row.comments]
        for i, value in enumerate(values, 1):
            wm.cell(r, i, _cell_value(value))
        r += 1
    _set_widths(wm, [28, 16, 18, 22, 22, 24, 32])

    wd = wb.create_sheet("Document Checklist")
    _header(wd, "DOCUMENT REQUEST & VALIDATION CHECKLIST", "G")
    for i, h in enumerate(
        ["Document", "Required", "Received?", "Not available", "File", "Uploaded", "Comments / Validation Notes"],
        1,
    ):
        cell = wd.cell(4, i, h)
        cell.font = _font(True, WHITE)
        cell.fill = _fill(GOLD)
    r = 5
    for doc in case.documents.select_related("document_type").prefetch_related("attachments"):
        files = list(doc.attachments.all())
        if not files and doc.file:
            files = [doc]
        if not files:
            wd.cell(r, 1, doc.document_type.name)
            wd.cell(r, 2, "Yes" if doc.document_type.is_required else "No")
            wd.cell(r, 3, doc.status_label())
            wd.cell(r, 4, "Yes" if doc.not_available else "No")
            wd.cell(r, 5, "")
            wd.cell(r, 6, "")
            wd.cell(r, 7, _cell_value(doc.comments))
            r += 1
            continue
        for att in files:
            name = att.display_name() if hasattr(att, "display_name") else (
                Path(att.file.name).name if att.file else ""
            )
            when = att.uploaded_at.strftime("%Y-%m-%d %H:%M") if getattr(att, "uploaded_at", None) else ""
            wd.cell(r, 1, doc.document_type.name)
            wd.cell(r, 2, "Yes" if doc.document_type.is_required else "No")
            wd.cell(r, 3, doc.status_label())
            wd.cell(r, 4, "Yes" if doc.not_available else "No")
            wd.cell(r, 5, _cell_value(name))
            wd.cell(r, 6, when)
            wd.cell(r, 7, _cell_value(doc.comments))
            r += 1
    _set_widths(wd, [42, 12, 14, 14, 36, 18, 50])

    wv = wb.create_sheet("Data Validation")
    _header(wv, "DATA QUALITY & VALIDATION CHECKLIST", "C")
    for i, h in enumerate(["Validation Item", "Status", "Comments / Follow-up"], 1):
        cell = wv.cell(4, i, h)
        cell.font = _font(True, WHITE)
        cell.fill = _fill(GOLD)
    r = 5
    for row in case.validations.select_related("item"):
        wv.cell(r, 1, _cell_value(row.item.name))
        wv.cell(r, 2, _cell_value(row.status))
        wv.cell(r, 3, _cell_value(row.comments))
        r += 1
    _set_widths(wv, [42, 18, 50])

    if include_market:
        orgs = list(orgs)
        analysis_note = (
            "Official is the midpoint of min/max (legacy payroll median only if both are blank). "
            "County P25/P50/P75 use that official figure plus staff forms for the job group; "
            "quartiles need at least two observations. "
            "Market percentiles are taken from comparator organisation medians (each organisation counts once). "
            "Ratio is county P50 ÷ market P50 (1.00 = at market). Scale uses market P25/P75 when two or more organisations have a figure."
        )
        analysis_rows = build_analysis_rows(case, orgs, case.survey_period, measure="basic")
        _write_job_analysis_sheet(
            wb,
            "Analysis Summary",
            "MARKET POSITION BY JOB GROUP (BASIC)",
            analysis_rows,
            orgs,
            analysis_note,
        )
        total_rows = build_analysis_rows(case, orgs, case.survey_period, measure="total")
        _write_job_analysis_sheet(
            wb,
            "Total Cash Position",
            "MARKET POSITION BY JOB GROUP (TOTAL CASH)",
            total_rows,
            orgs,
            analysis_note,
        )
        _write_rollup_sheet(
            wb,
            "By SRC Equivalent",
            "PAY BY SRC EQUIVALENT (BASIC)",
            rollup_rows(analysis_rows, "src_equivalent"),
        )
        _write_rollup_sheet(
            wb,
            "SRC Equivalent Total Cash",
            "PAY BY SRC EQUIVALENT (TOTAL CASH)",
            rollup_rows(total_rows, "src_equivalent"),
        )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_progress_workbook(cases) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Collection Progress"
    _header(ws, "COUNTY FILE COLLECTION PROGRESS", "H")
    heads = ["County", "Period", "Status", "Questionnaire %", "Jobs %", "Benefits %", "Documents %", "Certified"]
    for i, h in enumerate(heads, 1):
        cell = ws.cell(4, i, h)
        cell.font = _font(True, WHITE)
        cell.fill = _fill(GOLD)
    r = 5
    for case in cases:
        p = case.progress()
        values = [
            case.county.name,
            case.survey_period.name,
            case.get_status_display(),
            p["questionnaire"],
            p["jobs"],
            p["benefits"],
            p["documents"],
            "Yes" if p["certified"] else "No",
        ]
        for i, value in enumerate(values, 1):
            ws.cell(r, i, value)
        r += 1
    _set_widths(ws, [22, 22, 18, 16, 12, 12, 14, 12])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_questionnaire_extract(cases) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Questionnaire Extract"
    fields = [
        ("County", lambda c, q: c.county.name),
        ("Status", lambda c, q: c.get_status_display()),
        ("Department", lambda c, q: q.department),
        ("Respondent name", lambda c, q: q.respondent_name),
        ("Respondent designation", lambda c, q: q.respondent_designation),
        ("Respondent telephone", lambda c, q: q.respondent_telephone),
        ("Respondent email", lambda c, q: q.respondent_email),
        ("Permanent & pensionable", lambda c, q: q.permanent_pensionable),
        ("Contract", lambda c, q: q.contract_employees),
        ("Casual", lambda c, q: q.temporary_casual),
        ("Other staff", lambda c, q: q.other_employees),
        ("Total staff", lambda c, q: q.total_employees()),
        ("Approved policy", lambda c, q: q.approved_policy),
        ("Philosophy defined", lambda c, q: q.philosophy_defined),
        ("Market position", lambda c, q: q.market_position),
        ("Last review year", lambda c, q: q.last_review_year),
        ("Review triggers", lambda c, q: q.review_triggers),
        ("Number of grades", lambda c, q: q.number_of_grades),
        ("Progression factors", lambda c, q: q.progression_factors),
        ("If other, progression factors", lambda c, q: q.progression_factors_notes),
        ("Paid above maximum", lambda c, q: q.paid_above_maximum),
        ("If paid above maximum, why", lambda c, q: q.paid_above_maximum_why),
        ("Remunerative allowances", lambda c, q: q.remunerative_allowances),
        ("Non-remunerative allowances", lambda c, q: q.non_remunerative_allowances),
        ("Allowance expenditure", lambda c, q: q.annual_allowance_expenditure),
        ("Medical insurance", lambda c, q: q.medical_insurance),
        ("Medical cover arrangement", lambda c, q: q.medical_cover_how),
        ("Medical cover explanation", lambda c, q: q.medical_cover_notes),
        ("Medical cost / employee", lambda c, q: q.medical_cost_per_employee),
        ("Pension employer %", lambda c, q: q.pension_employer_pct),
        ("Pension employee %", lambda c, q: q.pension_employee_pct),
        ("Gratuity %", lambda c, q: q.gratuity_pct),
        ("Other significant benefits", lambda c, q: q.other_significant_benefits),
        ("Recruitment challenges", lambda c, q: q.recruitment_challenges),
        ("Hard-to-recruit families", lambda c, q: q.job_families_recruitment),
        ("Hard-to-retain families", lambda c, q: q.job_families_retention),
        ("Recruitment challenge reasons", lambda c, q: q.challenge_reasons),
        ("Retention challenge reasons", lambda c, q: q.retention_challenge_reasons),
        ("Benchmarking frequency", lambda c, q: q.benchmarking_frequency),
        ("Usual comparators", lambda c, q: q.usual_comparators),
        ("Internal equity", lambda c, q: q.internal_equity),
        ("Pay differences", lambda c, q: q.pay_differences),
        ("Pay difference causes", lambda c, q: q.pay_difference_causes),
        ("Basic % of cash", lambda c, q: q.basic_pct),
        ("Allowances % of cash", lambda c, q: q.allowances_pct),
        ("Other cash %", lambda c, q: q.other_cash_pct),
        ("Progression method", lambda c, q: q.progression_method),
        ("Average years in grade", lambda c, q: q.avg_years_in_grade),
        ("Placement factors", lambda c, q: q.placement_factors),
        ("Personnel % of budget", lambda c, q: q.personnel_pct_budget),
        ("Can afford ~15% increase", lambda c, q: q.can_accommodate_increase),
        ("Implementation approach", lambda c, q: q.implementation_approach),
        ("Implementation period", lambda c, q: q.implementation_period),
        ("Priority categories for increase", lambda c, q: q.priority_categories),
        ("Three challenges", lambda c, q: q.three_challenges),
        ("Three recommendations", lambda c, q: q.three_recommendations),
        ("Other comments", lambda c, q: q.other_comments),
    ]
    from catalog.models import QuestionItem

    extras = list(
        QuestionItem.objects.filter(is_builtin=False, is_active=True).order_by("section", "sort_order")
    )
    for item in extras:
        fields.append(
            (
                item.label,
                lambda c, q, question=item: (
                    q.extra_answers.filter(question=question).values_list("value", flat=True).first() or ""
                    if getattr(q, "pk", None)
                    else ""
                ),
            )
        )
    last_col = get_column_letter(len(fields))
    _header(ws, "CROSS-COUNTY QUESTIONNAIRE EXTRACT", last_col)
    _write_heads(ws, [h for h, _ in fields])
    r = 5
    from collection.models import Questionnaire as _Questionnaire

    for case in cases:
        q = case.get_questionnaire() or _Questionnaire(case=case)
        for i, (_, fn) in enumerate(fields, 1):
            value = fn(case, q)
            ws.cell(r, i, _cell_value(value))
        r += 1
    _set_widths(ws, [18] * len(fields))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _write_market_position_sheet(ws, heading, cases, orgs, period, measure="basic"):
    _header(ws, heading, "P")
    heads = [
        "County",
        "Status",
        "Job Group",
        "SRC Equivalent",
        "Official",
        "Staff n",
        "County P25",
        "County P50",
        "County P75",
        "Market n",
        "Market P25",
        "Market P50",
        "Market P75",
        "Ratio",
        "Position",
        "Scale",
    ]
    _write_heads(ws, heads)
    r = 5
    all_rows = []
    for case in cases:
        rows = build_analysis_rows(case, orgs, period, measure=measure)
        for item in rows:
            item["county_name"] = case.county.name
            item["county_status"] = case.get_status_display()
            all_rows.append(item)
            if not item["has_county"]:
                continue
            values = [
                item["county_name"],
                item["county_status"],
                item.get("job_group") or getattr(item.get("job"), "code", ""),
                item.get("src_equivalent") or "",
                _cell_value(item.get("official")),
                item.get("staff_n") or 0,
                _cell_value(item.get("county_p25")),
                _cell_value(item.get("county")),
                _cell_value(item.get("county_p75")),
                item.get("market_n") or 0,
                _cell_value(item.get("market_p25")),
                _cell_value(item.get("market")),
                _cell_value(item.get("market_p75")),
                _cell_value(item.get("position")),
                item.get("position_label") or "",
                item.get("scale_label") or "",
            ]
            for i, value in enumerate(values, 1):
                ws.cell(r, i, value)
            r += 1
    _set_widths(ws, [20, 16, 12, 16, 14, 10, 14, 14, 14, 12, 14, 14, 14, 10, 16, 22])
    return [item for item in all_rows if item.get("has_county")]


def build_analysis_workbook(cases, orgs, period) -> bytes:
    usable_orgs = orgs_with_pay(orgs, period) if period else list(orgs)
    wb = Workbook()
    ws = wb.active
    ws.title = "Market Position"
    paid_basic = _write_market_position_sheet(
        ws, "CROSS-COUNTY MARKET POSITION (BASIC)", cases, usable_orgs, period, "basic"
    )
    ws_total = wb.create_sheet("Total Cash Position")
    paid_total = _write_market_position_sheet(
        ws_total, "CROSS-COUNTY MARKET POSITION (TOTAL CASH)", cases, usable_orgs, period, "total"
    )
    _write_rollup_sheet(
        wb,
        "By SRC Equivalent",
        "PAY BY SRC EQUIVALENT (BASIC)",
        rollup_rows(paid_basic, "src_equivalent"),
    )
    _write_rollup_sheet(
        wb,
        "SRC Equivalent Total Cash",
        "PAY BY SRC EQUIVALENT (TOTAL CASH)",
        rollup_rows(paid_total, "src_equivalent"),
    )
    _write_comparator_columns_sheet(
        wb,
        "Comparator medians",
        "COMPARATOR ORGANISATION MEDIANS (BASIC)",
        paid_basic,
        usable_orgs,
    )
    _write_comparator_columns_sheet(
        wb,
        "Comparator total cash",
        "COMPARATOR ORGANISATION MEDIANS (TOTAL CASH)",
        paid_total,
        usable_orgs,
    )
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_individuals_extract(period, county_ids=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Individual responses"
    qs = IndividualResponse.objects.select_related("county", "job", "job__job_group").order_by(
        "county__name", "-submitted_at"
    )
    if period:
        qs = qs.filter(survey_period=period)
    if county_ids is not None:
        qs = qs.filter(county_id__in=list(county_ids))
    _fill_individuals_sheet(ws, "INDIVIDUAL STAFF SURVEY RESPONSES", qs)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_job_group_extract(cases) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Job Group Pay"
    _header(ws, "CROSS-COUNTY JOB GROUP PAY", "L")
    heads = [
        "County",
        "Status",
        "Job Group",
        "SRC Equivalent",
        "N/A",
        "Min basic",
        "Max basic",
        "Median (analysis)",
        "House",
        "Commuter",
        "Airtime",
        "Total of Other Allowances",
    ]
    _write_heads(ws, heads)
    r = 5
    from catalog.models import JobGroup

    for case in cases:
        for row in case.job_rows.select_related("job_group").filter(
            job_group__code__in=JobGroup.PAY_CODES
        ):
            values = [
                case.county.name,
                case.get_status_display(),
                row.job_group.title,
                _src_equivalent(row),
                "Yes" if row.not_applicable else "No",
                _cell_value(row.min_basic),
                _cell_value(row.max_basic),
                _cell_value(row.midpoint_basic()),
                _cell_value(row.house_allowance),
                _cell_value(row.commuter_allowance),
                _cell_value(row.airtime),
                _cell_value(row.other_allowances),
            ]
            for i, value in enumerate(values, 1):
                ws.cell(r, i, value)
            r += 1
    _set_widths(ws, [18, 14, 16, 16, 8, 14, 14, 16, 12, 12, 12, 22])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_benefits_extract(cases) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Benefits"
    _header(ws, "CROSS-COUNTY EMPLOYEE BENEFITS", "I")
    heads = [
        "County",
        "Status",
        "Benefit",
        "Provided?",
        "Interest charged?",
        "Employer contribution",
        "Employee contribution",
        "Eligibility criteria",
        "Explanatory comments",
    ]
    _write_heads(ws, heads)
    r = 5
    for case in cases:
        for row in case.benefit_rows.select_related("benefit").filter(benefit__is_active=True):
            values = [
                case.county.name,
                case.get_status_display(),
                row.benefit.name,
                row.provided,
                row.interest_charged,
                row.employer_contribution,
                row.employee_contribution,
                row.eligibility,
                row.comments,
            ]
            for i, value in enumerate(values, 1):
                ws.cell(r, i, _cell_value(value))
            r += 1
    _set_widths(ws, [18, 14, 28, 12, 16, 22, 22, 28, 32])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_documents_register(cases) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documents"
    _header(ws, "DOCUMENT REGISTER", "I")
    heads = [
        "County",
        "Status",
        "Document",
        "Required",
        "Received",
        "Not available",
        "File name",
        "Uploaded at",
        "Uploaded by",
        "Comments",
    ]
    _write_heads(ws, heads)
    r = 5
    for case in cases:
        for doc in case.documents.select_related("document_type", "uploaded_by").prefetch_related(
            "attachments", "attachments__uploaded_by"
        ):
            files = list(doc.attachments.all())
            if not files and doc.file:
                files = [doc]
            if not files:
                values = [
                    case.county.name,
                    case.get_status_display(),
                    doc.document_type.name,
                    "Yes" if doc.document_type.is_required else "No",
                    doc.status_label(),
                    "Yes" if doc.not_available else "No",
                    "",
                    "",
                    "",
                    doc.comments,
                ]
                for i, value in enumerate(values, 1):
                    ws.cell(r, i, value)
                r += 1
                continue
            for att in files:
                name = att.display_name() if hasattr(att, "display_name") else (
                    Path(att.file.name).name if att.file else ""
                )
                when = att.uploaded_at.strftime("%Y-%m-%d %H:%M") if getattr(att, "uploaded_at", None) else ""
                who = ""
                if getattr(att, "uploaded_by_id", None):
                    who = str(att.uploaded_by)
                elif doc.uploaded_by_id:
                    who = str(doc.uploaded_by)
                values = [
                    case.county.name,
                    case.get_status_display(),
                    doc.document_type.name,
                    "Yes" if doc.document_type.is_required else "No",
                    doc.status_label(),
                    "Yes" if doc.not_available else "No",
                    name,
                    when,
                    who,
                    doc.comments,
                ]
                for i, value in enumerate(values, 1):
                    ws.cell(r, i, value)
                r += 1
    _set_widths(ws, [18, 14, 40, 12, 12, 14, 36, 16, 22, 36])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _unique_zip_name(used, name):
    if name not in used:
        used.add(name)
        return name
    path = Path(name)
    stem, ext, parent = path.stem, path.suffix, str(path.parent)
    suffix = 2
    while True:
        candidate = f"{parent}/{stem}-{suffix}{ext}" if parent not in {".", ""} else f"{stem}-{suffix}{ext}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        suffix += 1


def build_county_complete_zip(case: CountyCase, include_market=True) -> bytes:
    """One folder: county workbook plus every uploaded supporting file."""
    ensure_case_structure(case)
    include_market = bool(include_market)
    stamp = timezone.now()
    root = slugify(f"{case.county.name} {case.survey_period.code}") or "county"
    root = f"DCollect-{root}-complete"
    used = set()
    missing = []
    not_available = []
    file_count = 0

    docs = list(
        case.documents.select_related("document_type")
        .prefetch_related("attachments")
        .filter(document_type__is_active=True)
        .order_by("document_type__sort_order", "document_type__name")
    )
    for doc in docs:
        if doc.not_available and not doc.has_files():
            not_available.append(doc.document_type.name)
        elif not doc.has_files():
            missing.append(doc.document_type.name)

    readme_lines = [
        f"DCollect complete file — {case.county.name}",
        f"Survey period: {case.survey_period.name}",
        f"Status: {case.get_status_display()}",
        f"Downloaded: {stamp.strftime('%d %b %Y %H:%M')}",
        "",
        "Open this folder to review everything the county submitted.",
        "",
        "1-County-submission.xlsx",
        "    Questionnaire, job groups, staff survey responses, benefits,",
        "    document checklist and validation.",
    ]
    if include_market:
        readme_lines.append("    Reviewer copy also includes comparator pay and market position.")
    readme_lines += [
        "",
        "2-Supporting-documents/",
        "    One sub-folder per document heading, with every file uploaded",
        "    under that heading.",
        "",
    ]
    if not_available:
        readme_lines.append("Marked not available:")
        readme_lines.extend(f"  - {name}" for name in not_available)
        readme_lines.append("")
    if missing:
        readme_lines.append("No file received:")
        readme_lines.extend(f"  - {name}" for name in missing)
        readme_lines.append("")

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        workbook = build_county_pack(case, include_market=include_market)
        archive.writestr(f"{root}/1-County-submission.xlsx", workbook)
        archive.writestr(f"{root}/README.txt", "\n".join(readme_lines) + "\n")
        for _doc, folder, name, handle in iter_case_document_files(case):
            zip_name = _unique_zip_name(used, f"{root}/2-Supporting-documents/{folder}/{name}")
            try:
                with handle.open("rb") as source:
                    archive.writestr(zip_name, source.read())
                file_count += 1
            except Exception:
                continue
        if file_count == 0:
            archive.writestr(
                f"{root}/2-Supporting-documents/No-files-uploaded.txt",
                "This county file has no supporting documents uploaded yet.\n",
            )
    return buffer.getvalue()

